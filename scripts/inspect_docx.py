from docx import Document

def inspect_tables(docx_path):
    doc = Document(docx_path)
    print(f"Total tables: {len(doc.tables)}\n")
    
    for i, table in enumerate(doc.tables):
        print(f"--- Table {i} ---")
        for r, row in enumerate(table.rows):
            row_data = []
            for cell in row.cells:
                text = cell.text.strip().replace('\n', ' ')
                row_data.append(text[:50] + "..." if len(text) > 50 else text)
            print(f"Row {r}: {row_data}")
            if r > 3: # Just print first few rows of each table to get the structure
                print("...")
                break
        print("\n")

if __name__ == "__main__":
    import sys
    inspect_tables(sys.argv[1])
