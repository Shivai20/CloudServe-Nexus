import docx

doc = docx.Document(r'FDE_Capstone_Complete/Capstone_Pack/02_Stage_Workbooks/Stage_2_PRD_Template.docx')
print(f"Total tables: {len(doc.tables)}")
for i, t in enumerate(doc.tables):
    header = [c.text.strip().replace('\n', ' ') for c in t.rows[0].cells]
    print(f"Table {i}: {len(t.rows)} rows x {len(t.columns)} cols | Header: {header}")
