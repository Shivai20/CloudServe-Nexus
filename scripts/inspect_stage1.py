import docx

doc = docx.Document(r'FDE_Capstone_Complete/Capstone_Pack/02_Stage_Workbooks/Stage_1_Discovery_Workbook.docx')

print(f"Total tables: {len(doc.tables)}")
for i, table in enumerate(doc.tables):
    print(f"\n--- TABLE {i} ({len(table.rows)} rows x {len(table.columns)} cols) ---")
    for r_idx, row in enumerate(table.rows):
        cells = [c.text.strip().replace('\n', ' ') for c in row.cells]
        print(f"Row {r_idx}: {cells}")
