from docx import Document
from datetime import datetime

def populate_prd(docx_path, output_path):
    doc = Document(docx_path)
    
    # Table 1: Metadata
    table1 = doc.tables[1]
    table1.cell(1, 1).text = "1.0"
    table1.cell(2, 1).text = "Lead FDE (Antigravity)"
    table1.cell(3, 1).text = datetime.now().strftime("%Y-%m-%d")
    table1.cell(4, 1).text = "Approved Draft"

    # Table 2: Problem Statement
    table2 = doc.tables[2]
    table2.cell(1, 0).text = (
        "CloudServe's customer support suffers from high resolution times (7.0 hours avg) and inconsistent "
        "CSAT (2.97/5.0), primarily because support agents manually answer repetitive queries "
        "(71.4% of tickets are already covered in existing documentation). Furthermore, language equity "
        "issues delay non-fluent English tickets by an additional 58 minutes on average. A highly accurate, "
        "deterministic AI routing and response system is needed to automatically resolve document-answerable "
        "tickets while instantly escalating complex or sensitive ('must_not_auto_respond') issues (17.4% of volume). "
        "This system must be strictly auditable and grounded in empirical data to ensure 100% compliance with business safety redlines."
    )

    # Table 3: User personas
    table3 = doc.tables[3]
    personas = [
        ("Customers raising tickets", "Fast, accurate, multilingual answers to their issues with minimal wait time.", "Wait 7.0 hours on average for a human to reply.", "Reduced time to first reply/resolution; increased CSAT."),
        ("Tier one support agents", "To focus on complex, high-value issues rather than manually copy-pasting from documentation.", "Spend time answering 71.4% of queries that are already documented.", "Reduction in manual ticket volume; handling only appropriate escalations."),
        ("Tier two and specialist agents", "Accurate, instantaneous escalations for sensitive ('must_not_auto_respond') queries with full context.", "Receive misrouted tickets or deal with massive backlogs.", "100% of PII, security, billing issues routed immediately to them without auto-responses."),
        ("The head of support", "Verifiable, auditable automation metrics to ensure safety, quality, and SLA savings (A1-A12).", "Lacks deterministic insight into automation decisions and potential hallucination risks.", "Generates automated evaluation reports reconciling 1:1 with input tickets.")
    ]
    for i, (persona, need, current, success) in enumerate(personas):
        # We start at row 1, 2, 3, 4
        table3.cell(i+1, 0).text = persona
        table3.cell(i+1, 1).text = need
        table3.cell(i+1, 2).text = current
        table3.cell(i+1, 3).text = success

    # Table 4: Functional Requirements
    table4 = doc.tables[4]
    frs = [
        ("FR-01", "The system shall normalize tickets from Email, Chat, Docs Comments, and Forum into a unified schema.", "Must", "500 ticket dataset reveals 4 distinct channels.", "Ingestion tests pass and schema validates."),
        ("FR-02", "The system shall classify intent across 22 classes + Urgency with numeric confidence.", "Must", "Need for calibrated intent detection.", "Classification module outputs valid intent & confidence."),
        ("FR-03", "The system shall retrieve authoritative documentation passages from the 29 KB articles.", "Must", "71.4% answerable from docs.", "Retrieval returns real doc IDs and chunk IDs."),
        ("FR-04", "The system shall route tickets deterministically based on confidence threshold and safety gates.", "Must", "17.4% redline tickets require mandatory escalation.", "Identical inputs produce identical routing decisions."),
        ("FR-05", "The system shall generate grounded responses citing specific retrieved doc passages.", "Must", "Need to avoid unverified LLM hallucinations.", "Factual claims in response directly link to cited text."),
        ("FR-06", "The system shall automatically block responses and trigger escalation upon PII, hallucination, or toxic tone.", "Must", "Governance standards & 17.4% redlines.", "Safety guardrails successfully trigger escalations."),
        ("FR-07", "The system shall generate a structured audit log entry matching 1:1 with tickets.", "Must", "Requirement A8 for decision audit trail.", "Total log entries reconcile exactly with tickets processed."),
        ("FR-08", "The system shall accept --input and --output via CLI and evaluate unattended without crashing.", "Must", "Requirement A9 (The Gate).", "CLI script runs validation set end-to-end successfully."),
        ("FR-09", "The system shall degrade gracefully on model timeout, API rate limit, or empty retrieval.", "Must", "Requirement A11 for Resilient Fault Tolerance.", "System continues pipeline without crashing on errors."),
        ("FR-10", "The system shall generate a comprehensive markdown/JSON evaluation report covering volume and metrics.", "Must", "Requirement A10 for Automated Metrics Report.", "Report is automatically created post-execution."),
        ("FR-11", "The system shall process non-fluent English queries accurately and equitably.", "Must", "Baseline non-fluent takes 58m longer with high misunderstanding risk.", "No significant degradation in classification/routing for non-fluent queries."),
        ("FR-12", "The system shall include a comprehensive pytest test suite covering all components.", "Must", "Requirement A12 for Pass-First Test Suite.", "python -m pytest tests/ -v passes all tests.")
    ]
    # Check if we need to add rows
    while len(table4.rows) < len(frs) + 1:
        table4.add_row()
        
    for i, (req_id, req, priority, evidence, acceptance) in enumerate(frs):
        row = table4.rows[i+1]
        row.cells[0].text = req_id
        row.cells[1].text = req
        row.cells[2].text = priority
        row.cells[3].text = evidence
        row.cells[4].text = acceptance

    # Table 5: Non-Functional Requirements
    table5 = doc.tables[5]
    nfrs = [
        ("NFR-01", "Latency", "System must process tickets efficiently and handle large evaluation sets within reasonable time limits.", "End-to-end pipeline execution time measurements."),
        ("NFR-02", "Availability/Resilience", "System must handle API rate limits and model timeouts without breaking the batch pipeline.", "Fault tolerance unit and integration tests (A11)."),
        ("NFR-03", "Accuracy", "Responses must only use facts present in the retrieved 29 KB articles.", "LLM-as-judge citation evaluation and manual spot checks (A6)."),
        ("NFR-04", "Security/Privacy", "System must automatically detect and escalate tickets containing PII.", "Safety guardrails testing on PII injection datasets (A7)."),
        ("NFR-05", "Auditability", "100% deterministic decision logging.", "Audit log reconciliation script ensuring 1:1 parity (A8)."),
        ("NFR-06", "Maintainability", "System must install and execute cleanly via documented README commands.", "Clean checkout verification on a fresh environment (A1).")
    ]
    while len(table5.rows) < len(nfrs) + 1:
        table5.add_row()
        
    for i, (req_id, cat, req, verification) in enumerate(nfrs):
        row = table5.rows[i+1]
        row.cells[0].text = req_id
        row.cells[1].text = cat
        row.cells[2].text = req
        row.cells[3].text = verification

    # Table 6: Out of scope
    table6 = doc.tables[6]
    oos = [
        ("Multi-turn conversations", "System acts as a one-shot routing/response mechanism based on initial ticket data.", "System expansion to support conversational state tracking."),
        ("Live API integrations for actions", "Goal is information retrieval and routing; taking destructive actions is unsafe for V1.", "Implementation of a safe execution sandbox or 'human-in-the-loop' approval step."),
        ("Generic Chit-Chat", "Support system must remain highly professional and domain-specific.", "Change in company policy to prioritize 'personality' over 'accuracy'.")
    ]
    while len(table6.rows) < len(oos) + 1:
        table6.add_row()
        
    for i, (not_building, why_not, change) in enumerate(oos):
        row = table6.rows[i+1]
        row.cells[0].text = not_building
        row.cells[1].text = why_not
        row.cells[2].text = change

    # Table 7: Assumptions
    table7 = doc.tables[7]
    assumps = [
        ("The 29 KB articles are accurate and up-to-date.", "They are the official knowledge base provided for the project.", "The system will generate incorrect answers grounded in faulty documentation.", "Monitoring user feedback (CSAT) and doc-update logs."),
        ("The 500 dev tickets are representative of production traffic.", "Provided by stakeholders as a baseline.", "Routing thresholds may be miscalibrated for real-world volume.", "Continuous evaluation of traffic post-deployment."),
        ("LLM models will perform deterministically with temperature 0.0.", "Standard behavior for major LLM providers.", "Auditability is compromised if identical inputs yield varying outputs.", "Automated regression testing of identical inputs.")
    ]
    while len(table7.rows) < len(assumps) + 1:
        table7.add_row()
        
    for i, (assump, why, false_impact, find_out) in enumerate(assumps):
        row = table7.rows[i+1]
        row.cells[0].text = assump
        row.cells[1].text = why
        row.cells[2].text = false_impact
        row.cells[3].text = find_out

    # Table 8: Metrics
    table8 = doc.tables[8]
    table8.cell(1, 2).text = ">= 60%" # FCR
    table8.cell(1, 3).text = "Automated Evaluation Report"
    table8.cell(1, 4).text = "FDE System"
    
    table8.cell(2, 2).text = "< 5 minutes (for AI responses)" # Time to first reply
    table8.cell(2, 3).text = "Automated Evaluation Report"
    table8.cell(2, 4).text = "FDE System"
    
    table8.cell(3, 2).text = ">= 4.0 / 5" # CSAT
    table8.cell(3, 3).text = "Simulated Feedback / Stakeholder Review"
    table8.cell(3, 4).text = "Quality Assurance Team"
    
    table8.cell(4, 0).text = "Escalation rate"
    table8.cell(4, 1).text = "56.2%"
    table8.cell(4, 2).text = "<= 40% (Only necessary escalations)"
    table8.cell(4, 3).text = "Automated Evaluation Report"
    table8.cell(4, 4).text = "FDE System"

    # Table 9: Open questions
    table9 = doc.tables[9]
    oq = [
        ("What specific embedding model should be used for the Vector Store?", "Impacts retrieval latency and accuracy.", "FDE Lead", "Start of Build Phase"),
        ("Are there API cost constraints for LLM usage?", "May necessitate model tiering strategies.", "Stakeholders", "Prior to Production deployment")
    ]
    while len(table9.rows) < len(oq) + 1:
        table9.add_row()
        
    for i, (q, why, owner, resolve) in enumerate(oq):
        row = table9.rows[i+1]
        row.cells[0].text = q
        row.cells[1].text = why
        row.cells[2].text = owner
        row.cells[3].text = resolve

    doc.save(output_path)
    print(f"Successfully saved PRD to {output_path}")

if __name__ == "__main__":
    import sys
    populate_prd(sys.argv[1], sys.argv[1])
