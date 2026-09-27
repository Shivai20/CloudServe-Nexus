from docx import Document

def populate_stage3(docx_path, output_path):
    doc = Document(docx_path)
    
    # Table 2: Specifications
    # FR-01: Ingestion
    # FR-02: Classification
    # FR-03: Retrieval
    # FR-04: Routing
    # FR-05: Generation
    # FR-06: Guardrails
    # FR-07: Audit
    # FR-08: Unattended
    # FR-09: Graceful degradation
    # FR-10: Report
    # FR-11: Language equity
    # FR-12: Tests
    table2 = doc.tables[2]
    specs = [
        ("FR-01", "Normalize Email/Chat/Docs/Forum to unified schema", "Raw JSON ticket", "Pydantic Ticket model", "Validates without error"),
        ("FR-02", "Classify intent (22 classes) and Urgency", "Ticket text", "Intent, Confidence [0-1], Urgency", "High accuracy on validation set"),
        ("FR-03", "Semantic retrieval from 29 KB articles", "Ticket text, Vector store", "List of Chunk IDs & text", "Returns real doc IDs"),
        ("FR-04", "Deterministic routing engine", "Confidence score, Safety flag", "Action: Respond or Escalate", "Identical input = identical output"),
        ("FR-05", "Grounded generation citing KB", "Ticket text, Retrieved chunks", "Final response with citations", "Zero hallucinations"),
        ("FR-06", "Safety guardrails for PII/Tone", "Ticket text, Final response", "Block flag", "100% block on PII"),
        ("FR-07", "Audit log generation", "Ticket ID, Routing Action", "Audit Log JSON entry", "1:1 reconciliation"),
        ("FR-08", "CLI unattended execution", "--input and --output paths", "Output JSON file", "Runs without crashing"),
        ("FR-09", "Fault tolerance", "Network timeouts/Empty retrieval", "Escalation Action", "Pipeline continues"),
        ("FR-10", "Automated evaluation report", "Audit logs, Ground truth", "Markdown metrics report", "Calculates FCR, CSAT, Latency"),
        ("FR-11", "Language equity", "Non-fluent English text", "Same as fluent English", "Latency delta < 5%"),
        ("FR-12", "Pytest suite", "Source code", "Test execution results", "All tests pass")
    ]
    
    while len(table2.rows) < len(specs) + 1:
        table2.add_row()
        
    for i, (req_id, spec, inputs, outputs, acc) in enumerate(specs):
        row = table2.rows[i+1]
        row.cells[0].text = req_id
        row.cells[1].text = spec
        row.cells[2].text = inputs
        row.cells[3].text = outputs
        row.cells[4].text = acc

    # Table 3 & 4: Prompt 1 (Classification)
    doc.tables[3].cell(1, 1).text = "Intent Classification"
    doc.tables[3].cell(2, 1).text = "Build"
    doc.tables[3].cell(3, 1).text = "FR-02"
    doc.tables[3].cell(4, 1).text = "1.0"
    
    doc.tables[4].cell(1, 0).text = (
        "You are an expert customer support classifier for CloudServe.\n"
        "Your task is to classify the user's ticket into exactly ONE of the 22 valid intents.\n"
        "Input Ticket:\n<TICKET>\n{ticket_text}\n</TICKET>\n\n"
        "Valid Intents: {valid_intents_list}\n\n"
        "Output JSON exactly matching this schema:\n"
        "{{\n  \"intent\": \"string\",\n  \"confidence\": \"float between 0 and 1\",\n  \"urgency\": \"Low|Medium|High\"\n}}"
    )

    # Table 5 & 6: Prompt 2 (Guardrails)
    doc.tables[5].cell(1, 1).text = "Safety Guardrails"
    doc.tables[5].cell(2, 1).text = "Build"
    doc.tables[5].cell(3, 1).text = "FR-06"
    doc.tables[5].cell(4, 1).text = "1.0"
    
    doc.tables[6].cell(1, 0).text = (
        "You are a strict security and compliance gatekeeper.\n"
        "Analyze the following support ticket to determine if it MUST NOT be auto-responded to.\n"
        "Flag as 'must_not_auto_respond' if it contains:\n"
        "1. PII or credit card numbers\n"
        "2. Security vulnerabilities\n"
        "3. Complex billing disputes\n"
        "4. Legal/compliance threats\n\n"
        "Input Ticket:\n<TICKET>\n{ticket_text}\n</TICKET>\n\n"
        "Output JSON exactly matching this schema:\n"
        "{{\n  \"is_safe\": \"boolean\",\n  \"reason\": \"string (if unsafe)\"\n}}"
    )

    # Table 7 & 8: Prompt 3 (Grounded Generation)
    doc.tables[7].cell(1, 1).text = "Grounded Generation"
    doc.tables[7].cell(2, 1).text = "Build"
    doc.tables[7].cell(3, 1).text = "FR-05"
    doc.tables[7].cell(4, 1).text = "1.0"
    
    doc.tables[8].cell(1, 0).text = (
        "You are a CloudServe support agent. Your goal is to resolve the user's ticket.\n"
        "CRITICAL RULE: You may ONLY use facts explicitly stated in the provided Knowledge Base chunks.\n"
        "If the answer is not in the chunks, you must state 'I do not have enough information to answer this.'\n\n"
        "User Ticket:\n<TICKET>\n{ticket_text}\n</TICKET>\n\n"
        "Knowledge Base Chunks:\n<CHUNKS>\n{retrieved_chunks}\n</CHUNKS>\n\n"
        "Output JSON exactly matching this schema:\n"
        "{{\n  \"answer\": \"string (markdown formatted)\",\n  \"citations\": [\"chunk_id_1\", \"chunk_id_2\"]\n}}"
    )

    # Table 11: Traceability
    table11 = doc.tables[11]
    
    trace_data = [
        ("FR-01", "Yes", "N/A (Code-based)", "test_ingestion_schema", "None"),
        ("FR-02", "Yes", "Intent Classification", "test_classification_accuracy", "None"),
        ("FR-03", "Yes", "N/A (Vector DB)", "test_retrieval_metrics", "Need embedding model selection"),
        ("FR-04", "Yes", "N/A (Logic Engine)", "test_deterministic_routing", "None"),
        ("FR-05", "Yes", "Grounded Generation", "test_grounded_generation", "None"),
        ("FR-06", "Yes", "Safety Guardrails", "test_guardrails_escalation", "None"),
        ("FR-07", "Yes", "N/A (Code-based)", "test_audit_log_1to1", "None"),
        ("FR-08", "Yes", "N/A (CLI Setup)", "test_cli_unattended", "None"),
        ("FR-09", "Yes", "N/A (Error Handling)", "test_fault_tolerance", "None"),
        ("FR-10", "Yes", "N/A (Report Gen)", "test_report_generation", "None"),
        ("FR-11", "Yes", "Classification/Gen", "test_language_equity", "Requires non-fluent test data"),
        ("FR-12", "Yes", "N/A (Testing)", "test_suite_execution", "None"),
    ]
    
    while len(table11.rows) < len(trace_data) + 1:
        table11.add_row()
        
    for i, (req_id, spec_written, prompts, test_case, gaps) in enumerate(trace_data):
        row = table11.rows[i+1]
        row.cells[0].text = req_id
        row.cells[1].text = spec_written
        row.cells[2].text = prompts
        row.cells[3].text = test_case
        row.cells[4].text = gaps

    doc.save(output_path)
    print(f"Successfully saved Prompt Library to {output_path}")

if __name__ == "__main__":
    import sys
    populate_stage3(sys.argv[1], sys.argv[1])
